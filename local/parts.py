from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Iterator, Mapping

from openpyxl import load_workbook

REQUIRED_COLUMNS: tuple[str, ...] = ("entry date", "quantity", "order unit", "pn")

@dataclass(frozen=True)
class PartRow:
    date: str
    quantity: float
    unit: str
    pn: str

def read_parts_text(path: str | Path) -> str:
    return "\n".join(_format_row(row) for row in read_parts(path))

def read_parts(path: str | Path) -> list[PartRow]:
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == ".csv":
        return list(_read_csv(source))
    if suffix == ".xlsx":
        return list(_read_xlsx(source))
    raise ValueError(f"unsupported parts file type: {suffix or '<none>'}")

def _read_csv(path: Path) -> Iterator[PartRow]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        if reader.fieldnames is None:
            raise ValueError("CSV file has no header row")
        _validate_headers(reader.fieldnames)
        for raw in reader:
            yield _row_from_mapping(_normalize_mapping(raw))

def _read_xlsx(path: Path) -> Iterator[PartRow]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration as exc:
            raise ValueError("XLSX file has no header row") from exc

        _validate_headers(header)
        index = {_normalize_header(name): i for i, name in enumerate(header)}

        for values in rows:
            raw = {
                name: values[i] if i < len(values) else None
                for name, i in index.items()
            }
            if all(value is None or str(value).strip() == "" for value in raw.values()):
                continue
            yield _row_from_mapping(raw)
    finally:
        workbook.close()

def _validate_headers(headers: tuple[object, ...] | list[object]) -> None:
    normalized = {_normalize_header(name) for name in headers}
    missing = [column for column in REQUIRED_COLUMNS if column not in normalized]
    if missing:
        raise ValueError(
            "spreadsheet is missing required columns: " + ", ".join(missing)
        )

def _normalize_header(value: object) -> str:
    return str(value).strip().lower() if value is not None else ""

def _normalize_mapping(raw: Mapping[str, object]) -> dict[str, object]:
    return {_normalize_header(key): value for key, value in raw.items()}

def _row_from_mapping(raw: Mapping[str, object]) -> PartRow:
    return PartRow(
        date=_format_date(_required(raw, "entry date")),
        quantity=_to_float(_required(raw, "quantity")),
        unit=str(_required(raw, "order unit")).strip(),
        pn=str(_required(raw, "pn")).strip(),
    )

def _required(raw: Mapping[str, object], key: str) -> object:
    value = raw.get(key)
    if value is None or str(value).strip() == "":
        raise ValueError(f"spreadsheet row is missing '{key}'")
    return value

def _format_date(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()

def _to_float(value: object) -> float:
    return float(str(value).strip().replace(",", ""))

def _format_row(row: PartRow) -> str:
    return (
        f"date={row.date}; quantity={row.quantity:g}; "
        f"unit={row.unit}; pn={row.pn}"
    )
