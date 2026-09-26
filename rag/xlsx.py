import csv
from pathlib import Path

from openpyxl import load_workbook

def convert(path: str, output: str | None = None) -> str:
    sheet = _first_sheet(path)
    target = _target(path, output)
    _write(target, sheet)
    return target

def convert_all(path: str, folder: str | None = None) -> list[str]:
    sheets = _all_sheets(path)
    dest = _folder(path, folder)
    return [_write_sheet(s, dest) for s in sheets]

def _first_sheet(path: str):
    return _all_sheets(path)[0]

def _all_sheets(path: str):
    return load_workbook(path, data_only=True).worksheets

def _write(target: str, sheet) -> None:
    with open(target, "w", newline="", encoding="utf-8") as f:
        _dump(f, sheet)

def _write_sheet(sheet, folder: str) -> str:
    target = str(Path(folder) / f"{sheet.title}.csv")
    _write(target, sheet)
    return target

def _dump(handle, sheet) -> None:
    writer = csv.writer(handle)
    for row in sheet.iter_rows(values_only=True):
        writer.writerow(_clean(row))

def _clean(row) -> list[str]:
    return ["" if c is None else str(c) for c in row]

def _target(path: str, output: str | None) -> str:
    return output or str(Path(path).with_suffix(".csv"))

def _folder(path: str, folder: str | None) -> str:
    dest = Path(folder or Path(path).with_suffix(""))
    dest.mkdir(parents=True, exist_ok=True)
    return str(dest)
