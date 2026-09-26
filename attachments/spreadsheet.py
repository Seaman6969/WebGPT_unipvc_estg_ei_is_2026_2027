import csv

from openpyxl import load_workbook

def read_xlsx(path: str) -> str:
    book = load_workbook(path, data_only=True)
    return _join_sheets(book.worksheets)

def _join_sheets(sheets) -> str:
    return "\n\n".join(_sheet(s) for s in sheets)

def _sheet(sheet) -> str:
    rows = [_row(r) for r in sheet.iter_rows(values_only=True)]
    return f"[{sheet.title}]\n" + "\n".join(rows)

def _row(row) -> str:
    return " | ".join(_cell(c) for c in row)

def _cell(c) -> str:
    return "" if c is None else str(c)

def read_csv(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return _rows(f)

def _rows(handle) -> str:
    return "\n".join(" | ".join(row) for row in csv.reader(handle))
